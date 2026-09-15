#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
역기획 문서 렌더러 (reverse-spec / reverse-prd 공용)

Claude가 Step 1~3에서 구성한 Markdown 본문을 받아 PDF 또는 DOCX로 출력한다.
스킬 본문에 Python을 인라인하지 않고 이 스크립트를 호출하므로,
allowed-tools 를 `Bash(python ${CLAUDE_SKILL_DIR}/scripts/*)` 로 좁힐 수 있다.

사용법:
  python render.py --input doc.md --title "역기획 PRD" --outdir reverse-prd-output          # PDF+HTML 세트(기본)
  python render.py --input doc.md --format docx,html --title "역기획 정책서" --outdir out   # DOCX+HTML

옵션:
  --input   구성된 Markdown 본문 파일 경로 (필수)
  --format  쉼표 구분 다중 지정: pdf,html(기본) | docx | html 조합
  --title   문서 머리말/제목          (기본 "역기획 문서")
  --accent  강조색 hex                (기본 #1f4e79)
  --outdir  출력 디렉토리             (기본 reverse-spec-output)
  --name    파일명 접두어             (기본 reverse_spec)
  --lang    HTML/PDF lang 속성        (기본 ko; 영문 문서는 en)
"""
import argparse
import datetime
import pathlib
import re
import sys


import urllib.parse

# PDF 렌더링 중 실제로 외부에서 가져올 필요가 있는 리소스는 구글 폰트뿐이다.
# 그 외 호스트/스킴을 전부 막아 SSRF(내부망·클라우드 메타데이터 주소 접근) 표면을
# "file:// 차단"보다 넓게 차단한다 — 분석 대상 코드에서 뽑아낸 문구가 이스케이프
# 없이 본문에 섞여 `<img src="http://169.254.169.254/...">` 같은 태그가 되더라도
# 이 fetcher가 요청 자체를 거부한다.
_ALLOWED_FETCH_HOSTS = {"fonts.googleapis.com", "fonts.gstatic.com"}


def _pdf_url_fetcher(url: str, *args, **kwargs):
    """weasyprint용 URL fetcher.

    보안 검증 발견 (1차): 기본 fetcher가 file:// 스킴을 그대로 가져와, 분석 대상이
    통제 못하는 문서에 `<link href="file:///etc/passwd">` 같은 참조가 있으면 로컬
    파일을 읽어 PDF에 임베드할 수 있었다(정보 노출).
    보안 검증 발견 (2차, 재점검): file:// 만 막고 http(s)://는 전부 허용하고 있어,
    SSRF로 내부망/클라우드 메타데이터 엔드포인트(예: http://169.254.169.254/...)에
    접근할 수 있는 여지가 남아 있었다. https + 알려진 폰트 호스트만 허용하는
    화이트리스트로 좁힌다(그 외 스킴·호스트는 전부 거부)."""
    parsed = urllib.parse.urlsplit(url)
    if parsed.scheme != "https" or parsed.hostname not in _ALLOWED_FETCH_HOSTS:
        raise ValueError(f"보안: 허용되지 않은 외부 리소스 참조 차단됨 ({url[:60]})")
    from weasyprint.urls import default_url_fetcher
    return default_url_fetcher(url, *args, **kwargs)


# 보안 검증 발견(재검토, PR 리뷰): 여는/닫는 태그를 통째로 한 쌍으로 매칭하려던
# 이전 버전(`<script>...내용...</script>` 전체를 `.*?`로 탐색)은, 닫히지 않는
# `<script>` 토큰이 반복되는 적대적 입력에서 각 시작 위치마다 나머지 문서
# 전체를 상대로 백트래킹해 초선형(superlinear) 시간이 걸렸다 — 1,000개 반복
# (~8KB)만으로 10초 이상 걸리는 CPU DoS가 실측 재현됨. 여는/닫는 태그를 각각
# 개별적으로(내용은 남기고 태그만) 제거하는 방식으로 바꿔 매 매치가 다음 `>`
# 하나로 폭이 정해지는 선형 시간 패턴으로 교체했다 — 내용이 페이지에 텍스트로
# 남더라도 실행 가능한 태그 자체가 사라지므로 목적(스크립트 실행 차단)은
# 동일하게 달성된다.
_DANGEROUS_TAGS = re.compile(
    r"<\s*/?\s*(?:script|iframe|object|embed|style|link)\b[^>]*>", re.I)
# 진짜 태그(<...>) 구간에서만 이벤트 속성/위험 스킴을 다듬는다 — 코드스팬 등
# 이미 HTML 엔티티로 이스케이프된 텍스트는 리터럴 `<`/`>`가 없어 이 패턴에
# 매칭되지 않으므로 건드리지 않는다.
_TAG = re.compile(r"<[a-zA-Z][^<>]*>")
_EVENT_ATTR = re.compile(r'\son\w+\s*=\s*("[^"]*"|\'[^\']*\'|[^\s>]+)', re.I)
# 보안 검증 발견(재검토, PR 리뷰): 값이 따옴표로 감싸여 있을 때만 매칭해,
# `href=javascript:alert(1)`처럼 따옴표 없는 속성값은 그대로 통과시켰다 —
# 따옴표는 선택 사항으로 바꿔 두 형태 모두 잡는다.
_DANGEROUS_SCHEME_ATTR = re.compile(
    r'\b(href|src)(\s*=\s*)(["\']?)\s*(javascript:|vbscript:|data:text/html)', re.I)


def _clean_tag(m: "re.Match") -> str:
    tag = m.group(0)
    tag = _EVENT_ATTR.sub("", tag)
    tag = _DANGEROUS_SCHEME_ATTR.sub(r"\1\2\3blocked:", tag)
    return tag


def _sanitize_html_fragment(html: str) -> str:
    """python-markdown은 raw HTML을 기본적으로 그대로 통과시킨다(safe_mode 없음).
    표 셀 텍스트는 extract.py의 _escape_cell이 이미 이스케이프하지만, Step 2/3에서
    LLM이 메시지 원문을 표가 아닌 본문 서술로 옮겨 적으면 그 경로는 보호되지
    않는다 — 렌더링 최종 단계의 방어선(defense-in-depth)으로 실행 가능한 태그·
    이벤트 핸들러·스크립트성 URL 스킴을 무력화한다.

    처음 구현에서 이벤트 속성 정규식을 문서 전체에 바로 적용했더니, 코드스팬
    안의 (이미 이스케이프된) 예시 문구 뒤에 오는 `&gt;` 같은 엔티티까지 값으로
    먹어치워 원문을 훼손했다(재검증 발견) — 실제 태그(`<...>`) 구간으로만
    범위를 좁혀 텍스트 콘텐츠는 절대 건드리지 않도록 고쳤다."""
    html = _DANGEROUS_TAGS.sub("", html)
    html = _TAG.sub(_clean_tag, html)
    return html


def build_html(md_text: str, title: str, accent: str, lang: str = "ko") -> str:
    import markdown
    # python-markdown 코어에 취소선(~~text~~)이 없어 <del>로 선치환 (HTML/PDF 공통)
    md_text = re.sub(r"~~(.+?)~~", r"<del>\1</del>", md_text)
    body = markdown.markdown(md_text, extensions=["tables", "toc", "fenced_code"])
    body = _sanitize_html_fragment(body)
    return f"""<!DOCTYPE html>
<html lang="{lang}">
<head>
<meta charset="UTF-8">
<style>
  @import url('https://fonts.googleapis.com/css2?family=Noto+Sans+KR:wght@400;700&display=swap');
  body {{ font-family: 'Noto Sans KR', sans-serif; font-size: 11pt; line-height: 1.8;
          color: #1a1a1a; margin: 0; padding: 0; }}
  h1 {{ font-size: 20pt; border-bottom: 2px solid {accent}; padding-bottom: 8px;
        margin-top: 40px; color: {accent}; }}
  h2 {{ font-size: 15pt; border-left: 4px solid {accent}; padding-left: 10px; margin-top: 30px; }}
  h3 {{ font-size: 12pt; color: #444; margin-top: 20px; }}
  table {{ width: 100%; border-collapse: collapse; margin: 16px 0; font-size: 10pt; }}
  th {{ background: {accent}; color: white; padding: 8px 12px; text-align: left; }}
  td {{ border: 1px solid #ccc; padding: 7px 12px; vertical-align: top; }}
  tr:nth-child(even) td {{ background: #f5f7fa; }}
  code {{ background: #eef2f7; padding: 2px 5px; border-radius: 3px; font-size: 9.5pt; }}
  pre {{ background: #1e1e1e; color: #d4d4d4; padding: 14px; border-radius: 6px;
         font-size: 9pt; overflow-x: auto; }}
  pre code {{ background: transparent; color: inherit; padding: 0; }}
  blockquote {{ border-left: 4px solid #c47a00; background:#fff8ec; margin:16px 0;
               padding:8px 16px; color:#7a5200; }}
  /* 브라우저 화면 전용 여백 — weasyprint(PDF)는 print 매체라 무시하고 @page 여백을 쓴다 */
  @media screen {{
    body {{ max-width: 860px; margin: 0 auto; padding: 32px 24px; }}
  }}
  @page {{
    size: A4;
    margin: 20mm 18mm 20mm 20mm;
    @top-center {{ content: "{title}"; font-size: 9pt; color: #888; }}
    @bottom-right {{ content: counter(page) " / " counter(pages); font-size: 9pt; color: #888; }}
  }}
</style>
</head>
<body>
{body}
</body>
</html>"""


def render_pdf(md_text: str, out_path: pathlib.Path, title: str, accent: str, lang: str) -> None:
    from weasyprint import HTML
    HTML(string=build_html(md_text, title, accent, lang),
         url_fetcher=_pdf_url_fetcher).write_pdf(str(out_path))


def render_html(md_text: str, out_path: pathlib.Path, title: str, accent: str, lang: str) -> None:
    # 미리보기/디버그용 정적 HTML (A4 @page 규칙은 브라우저에서 무시됨)
    out_path.write_text(build_html(md_text, title, accent, lang), encoding="utf-8")


# 인라인 마크다운 → 평문 변환 규칙 (순서 중요: 이미지→링크 순).
_IMG = re.compile(r"!\[([^\]]*)\]\([^)]*\)")           # ![alt](url) → alt
_LINK = re.compile(r"\[([^\]]+)\]\([^)]*\)")            # [text](url) → text
_BOLD = re.compile(r"\*\*(.+?)\*\*")
_STRIKE = re.compile(r"~~(.+?)~~")
_CODE = re.compile(r"`+([^`]+)`+")                      # `code` / ``code`` 모두


def _clean_inline(text: str) -> str:
    """docx 평문용: 인라인 마크다운 기호(**bold**, `code`, [링크](url),
    ![이미지](url), ~~취소선~~)를 제거하고, 이스케이프된 파이프(\\|)와
    _escape_cell이 넣은 HTML 엔티티(&lt; 등)를 리터럴로 되돌린다.
    (렌더링 검증 발견: 링크/이미지/취소선이 원문 기호 그대로 노출되던 문제 수정.)"""
    text = _IMG.sub(r"\1", text)
    text = _LINK.sub(r"\1", text)
    text = _BOLD.sub(r"\1", text)
    text = _STRIKE.sub(r"\1", text)
    text = _CODE.sub(r"\1", text)
    text = text.replace("<br/>", "\n").replace("<br>", "\n")
    text = text.replace("\\|", "|")
    # _escape_cell의 HTML 엔티티를 Word용 리터럴로 복원 (&amp; 는 마지막에)
    text = text.replace("&lt;", "<").replace("&gt;", ">").replace("&amp;", "&")
    return text


def _split_table_row(line: str) -> list:
    """GFM 파이프 테이블 행을 셀로 분리한다. 백틱(코드 스팬) 안의 파이프와
    백슬래시로 이스케이프된 파이프(\\|)는 구분자로 보지 않는다 — 단순
    line.split("|")는 이 두 경우를 무시해 JS `a || b` 같은 조건문이 들어간
    셀에서 뒤 컬럼이 사라지는 버그가 있었다 (실전 검증 발견).
    렌더링 검증 추가 발견: 백틱을 단순 토글로 처리하면 이중 백틱(``code``)에서
    깨졌으므로, 연속 백틱 런(run) 길이가 같은 쌍으로 코드 스팬을 닫는다."""
    cells, buf, i = [], [], 0
    backtick_run = 0  # 0이면 코드스팬 밖, >0이면 그 길이의 런으로 열린 상태
    while i < len(line):
        ch = line[i]
        if ch == "\\" and i + 1 < len(line) and line[i + 1] == "|":
            buf.append("|")
            i += 2
            continue
        if ch == "`":
            j = i
            while j < len(line) and line[j] == "`":
                j += 1
            run = j - i
            if backtick_run == 0:
                backtick_run = run           # 코드 스팬 열기
            elif backtick_run == run:
                backtick_run = 0             # 같은 길이 런으로 닫기
            buf.append(line[i:j])
            i = j
            continue
        if ch == "|" and backtick_run == 0:
            cells.append("".join(buf))
            buf = []
        else:
            buf.append(ch)
        i += 1
    cells.append("".join(buf))
    # 라인이 "| a | b |" 형태면 앞뒤 빈 셀(선행/후행 |)이 생기므로 제거
    stripped = line.strip()
    if stripped.startswith("|") and cells and cells[0].strip() == "":
        cells = cells[1:]
    if stripped.endswith("|") and cells and cells[-1].strip() == "":
        cells = cells[:-1]
    return [c.strip() for c in cells]


def _add_list_para(doc, text: str, base_style: str, nest: int):
    """중첩 레벨(nest)에 맞는 리스트 스타일을 적용하되, 해당 스타일이 문서
    템플릿에 없으면 기본 스타일로 안전 폴백한다 ('List Bullet 2' 등이 없는
    템플릿에서 KeyError로 죽지 않도록)."""
    style = base_style if nest == 0 else f"{base_style} {nest + 1}"
    try:
        doc.add_paragraph(text, style=style)
    except KeyError:
        doc.add_paragraph(text, style=base_style)


def render_docx(md_text: str, out_path: pathlib.Path, title: str) -> None:
    from docx import Document
    from docx.shared import Cm

    doc = Document()
    doc.core_properties.title = title
    sec = doc.sections[0]
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.0)

    # PDF의 @top-center 머리말과 동일하게 페이지 헤더에 문서 제목 표시
    header_p = sec.header.paragraphs[0]
    header_p.text = title
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    lines = md_text.splitlines()
    i = 0
    in_code = False
    code_buf: list[str] = []
    table_buf: list[str] = []

    def flush_table() -> None:
        # GFM 파이프 테이블: 2번째 행이 |---| 구분선
        rows = [r for r in table_buf if r.strip()]
        if len(rows) >= 2:
            header = _split_table_row(rows[0])
            data = [_split_table_row(r) for r in rows[2:]]
            t = doc.add_table(rows=1 + len(data), cols=len(header))
            t.style = "Table Grid"
            for j, h in enumerate(header):
                cell = t.rows[0].cells[j]
                cell.text = _clean_inline(h)
                if cell.paragraphs[0].runs:
                    cell.paragraphs[0].runs[0].bold = True
            for r_idx, row in enumerate(data):
                for j in range(len(header)):
                    val = row[j] if j < len(row) else ""
                    t.rows[r_idx + 1].cells[j].text = _clean_inline(val)
        table_buf.clear()

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # 코드 펜스
        if stripped.startswith("```"):
            if in_code:
                doc.add_paragraph("\n".join(code_buf), style="Intense Quote")
                code_buf.clear()
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        # 테이블 누적
        if stripped.startswith("|"):
            table_buf.append(line)
            i += 1
            continue
        elif table_buf:
            flush_table()

        if not stripped:
            i += 1
            continue

        # 헤딩 (H1~H6 — 렌더링 검증 발견: #{1,4}라 #####/###### 가 문단으로 강등됐음)
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            doc.add_heading(_clean_inline(m.group(2)), level=min(level, 6))
            i += 1
            continue

        # 인용
        if stripped.startswith(">"):
            p = doc.add_paragraph(_clean_inline(stripped.lstrip("> ").rstrip()))
            p.style = "Quote"
            i += 1
            continue

        # 들여쓰기 기반 중첩 리스트 레벨 (렌더링 검증 발견: strip 후 검사해 중첩
        # 정보가 사라지고 모두 같은 레벨로 평탄화됐음 — 원본 라인의 선행 공백으로
        # 레벨을 계산한다). 공백 2칸 = 1레벨.
        indent = len(line) - len(line.lstrip(" "))
        nest = min(indent // 2, 2)  # docx 기본 List 스타일은 3레벨까지

        # 번호 목록 (1. 2. …) — 렌더링 검증 발견: 어떤 분기에도 안 걸려 문단으로
        # 강등되며 Word 자동번호가 사라졌음
        mo = re.match(r"^\d+\.\s+(.*)$", stripped)
        if mo:
            _add_list_para(doc, _clean_inline(mo.group(1)), "List Number", nest)
            i += 1
            continue

        # 불릿/체크리스트
        if re.match(r"^[-*]\s+", stripped):
            text = re.sub(r"^[-*]\s+", "", stripped)
            _add_list_para(doc, _clean_inline(text), "List Bullet", nest)
            i += 1
            continue

        # 일반 문단
        doc.add_paragraph(_clean_inline(stripped))
        i += 1

    if table_buf:
        flush_table()

    doc.save(str(out_path))


def main() -> int:
    ap = argparse.ArgumentParser(description="역기획 문서 렌더러 (pdf/docx/html)")
    ap.add_argument("--input", required=True, help="구성된 Markdown 본문 파일")
    ap.add_argument("--format", default="pdf,html",
                    help="쉼표 구분 다중 지정 가능: pdf,html(기본) / docx,html / pdf 등")
    ap.add_argument("--title", default="역기획 문서")
    ap.add_argument("--accent", default="#1f4e79")
    ap.add_argument("--outdir", default="reverse-spec-output")
    ap.add_argument("--name", default="reverse_spec")
    ap.add_argument("--lang", default="ko", help="HTML/PDF lang 속성 (ko|en 등)")
    args = ap.parse_args()

    formats = [f.strip() for f in args.format.split(",") if f.strip()]
    invalid = [f for f in formats if f not in ("pdf", "docx", "html")]
    if invalid:
        print(f"❌ 지원하지 않는 형식: {', '.join(invalid)} (pdf|docx|html)", file=sys.stderr)
        return 1

    md_text = pathlib.Path(args.input).read_text(encoding="utf-8")
    out_dir = pathlib.Path(args.outdir)
    out_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")

    # 보안: --name은 "reverse_prd" 같은 파일명 접두어여야 하는데, 경로 구분자나
    # ".."가 섞여 들어오면 out_dir 밖에 파일을 쓰는 경로 조작이 될 수 있다
    # (이 스크립트는 에이전트가 분석 대상 코드/문서 내용을 바탕으로 인자를 구성해
    # 호출하므로, 신뢰 못 할 입력이 --name까지 흘러들 가능성을 전제로 방어한다).
    # 디렉터리 구성요소를 제거해 순수 파일명만 남긴다.
    safe_name = pathlib.PurePosixPath(args.name.replace("\\", "/")).name or "reverse_doc"
    if safe_name != args.name:
        print(f"⚠️  --name 값을 안전한 파일명으로 정규화함: {args.name!r} → {safe_name!r}",
              file=sys.stderr)

    for fmt in formats:  # 같은 타임스탬프로 세트 생성 (예: _104205.pdf + _104205.html)
        out_path = out_dir / f"{safe_name}_{ts}.{fmt}"
        if fmt == "pdf":
            render_pdf(md_text, out_path, args.title, args.accent, args.lang)
        elif fmt == "docx":
            render_docx(md_text, out_path, args.title)
        else:
            render_html(md_text, out_path, args.title, args.accent, args.lang)
        print(f"✅ {fmt.upper()} 생성 완료: {out_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
